# -*- coding: utf-8 -*-
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

DOC_PATH = 'TCC_HILTON_DOS_SANTOS_HOMEM 2026-2.revSimonedocx.docx'
doc = docx.Document(DOC_PATH)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def insert_code_box_before(p_target, title_text, code_str, note_text):
    # 1. Title / Caption
    p_title = p_target.insert_paragraph_before()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.line_spacing = 1.15
    p_title.paragraph_format.space_before = Pt(10)
    p_title.paragraph_format.space_after = Pt(4)
    p_title.paragraph_format.first_line_indent = Inches(0)
    r_t = p_title.add_run(title_text)
    r_t.font.name = 'Arial'
    r_t.font.size = Pt(10.5)
    r_t.font.bold = True
    r_t.font.color.rgb = RGBColor(0, 0, 0)
    
    # 2. Table containing code
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, '0F172A') # Slate 900 (dark mode)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    cell.paragraphs[0].text = ''
    lines = code_str.strip().splitlines()
    for l_idx, line in enumerate(lines):
        p_c = cell.paragraphs[0] if l_idx == 0 else cell.add_paragraph()
        p_c.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_c.paragraph_format.line_spacing = 1.05
        p_c.paragraph_format.space_before = Pt(0)
        p_c.paragraph_format.space_after = Pt(2)
        p_c.paragraph_format.first_line_indent = Inches(0)
        
        run_c = p_c.add_run(line)
        run_c.font.name = 'Consolas'
        run_c.font.size = Pt(9.5)
        run_c.font.color.rgb = RGBColor(226, 232, 240) # Slate 200
        
    col_w = Inches(6.5)
    table.rows[0].cells[0].width = col_w
    p_target._p.addprevious(table._tbl)
    
    # 3. Source / Note
    p_source = p_target.insert_paragraph_before()
    p_source.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_source.paragraph_format.line_spacing = 1.15
    p_source.paragraph_format.space_before = Pt(3)
    p_source.paragraph_format.space_after = Pt(10)
    p_source.paragraph_format.first_line_indent = Inches(0)
    r_s = p_source.add_run(f'Fonte: Elaborado pelo autor a partir do código-fonte ({note_text}).')
    r_s.font.name = 'Arial'
    r_s.font.size = Pt(9.5)
    r_s.font.italic = True
    r_s.font.color.rgb = RGBColor(71, 85, 105)

# Code snippets to insert:

# SNIPPET 1: Ingestão CSV (Seção 5.3)
code_1 = """// Leitura assíncrona, sanitização e mapeamento do plantel via CSV
const handleImportAtletasCSV = (e) => {
  const file = e.target.files[0];
  if (!file) return;
  const reader = new FileReader();
  reader.onload = (evt) => {
    const text = evt.target.result;
    const lines = text.split(/\\r?\\n/).map(l => l.trim()).filter(l => l.length > 0);
    if (lines.length < 2) return;
    const newAtletas = [];
    for (let i = 1; i < lines.length; i++) {
      const parts = lines[i].split(/[;,]/).map(p => p.trim());
      if (parts.length >= 6) {
        newAtletas.push({
          id: i,
          name: parts[0] || `Atleta ${i}`,
          min: parseInt(parts[1]) || 0,
          max: parseInt(parts[2]) || 3,
          sex: parts[3] || 'M',
          age: parseInt(parts[4]) || 30,
          score: parseFloat(parts[5].replace(',', '.')) || 0.5
        });
      }
    }
    if (newAtletas.length > 0) setAtletas(newAtletas);
  };
  reader.readAsText(file);
};"""

# SNIPPET 2: Regras Manuais (Seção 5.4)
code_2 = """// Filtragem de candidatos elegíveis respeitando regras manuais de bloqueio (RF06)
const candidates = inputAtletas.filter(atleta => {
  // 1. Não pode estar previamente alocado na tripulação
  if (crew.some(c => c.id === atleta.id)) return false;

  // 2. Limite máximo individual de participações (Equação 3.5)
  if (athleteCounts[atleta.id] >= atleta.max) return false;

  // 3. Janela fisiológica de descanso (Equação 3.7): intervalo eta
  const assigned = athleteAssignedProvas[atleta.id] || [];
  const restConflict = assigned.some(idx => Math.abs(idx - pIndex) <= inputConfig.restInterval);
  if (restConflict) return false;

  // 4. Aplicação rigorosa das regras de bloqueio cadastradas pelo técnico
  const blockRules = inputRestricoes.filter(r => r.type === 'block' && r.atletaId === atleta.id);
  for (const bRule of blockRules) {
    if (bRule.blockedBoats?.includes(prova.boat)) return false;
    if (bRule.blockedProvaUids?.includes(prova.uid)) return false;
  }
  return true;
});"""

# SNIPPET 3: Validação Pré-Voo (Seção 5.5)
code_3 = """// Auditoria diagnóstica pré-voo em 4 dimensões antes de autorizar o solver
const getValidationStatus = () => {
  const configValid = config.penalty >= 0 && config.restInterval >= 0;
  const provasValid = provas.length > 0 && provas.every(p => 
    p.id.trim() !== '' && p.date.trim() !== '' && p.time.trim() !== '');
  const atletasValid = atletas.length > 0 && atletas.every(a => 
    a.name.trim() !== '' && a.min <= a.max && a.min >= 0);
  const restricoesValid = restricoes.every(r => {
    if (r.type === 'block') {
      const atletaExists = atletas.some(a => a.id === r.atletaId);
      const validProvas = (r.blockedProvaUids || []).every(uid => provas.some(p => p.uid === uid));
      return atletaExists && validProvas;
    } else {
      const provaExists = provas.some(p => p.uid === r.provaUid);
      const validAtletas = (r.atletaIds || []).length > 0 && 
        (r.atletaIds || []).every(id => atletas.some(a => a.id === id));
      return provaExists && validAtletas;
    }
  });
  return { configValid, provasValid, atletasValid, restricoesValid };
};"""

# SNIPPET 4: Otimizador e Penalização Etária (Seção 5.5)
code_4 = """// Cálculo do grau de desempenho individual com penalização etária (Equação 3.8)
const getAdjustedScore = (atleta) => {
  const ageDiff = atleta.age - targetMinAge;
  return atleta.score - (ageDiff * inputConfig.penalty);
};

// Formação da tripulação e cálculo do grau final médio do barco (b_jk)
if (crew.length === seats) {
  const avgAge = (crew.reduce((sum, a) => sum + a.age, 0) / seats).toFixed(1);
  const adjustedScores = crew.map(a => {
    const ageDiff = a.age - targetMinAge;
    return Math.max(0, a.score - (ageDiff * inputConfig.penalty));
  });
  const boatScore = (adjustedScores.reduce((sum, s) => sum + s, 0) / seats);
  // Registro da guarnição otimizada para o relatório final
  allocatedProvas.push({ id: prova.id, boat: prova.boat, class: prova.ageClass, age: avgAge, score: boatScore.toFixed(3) });
}"""

# Find paragraphs where to insert
p_5_4_header = None
p_5_5_header = None
p_5_5_result_ph = None
p_5_6_header = None

for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t.startswith('5.4 Motor de Regras'):
        p_5_4_header = p
    elif t.startswith('5.5 Validação Pré-Voo') or t.startswith('5.5 Validao Pr-Voo'):
        p_5_5_header = p
    elif t.startswith('[insira aqui foto de tela de Valida') or t.startswith('[insira aqui foto de tela de Validao'):
        p_5_5_result_ph = p
    elif t.startswith('5.6 Central de Ajuda'):
        p_5_6_header = p

print("Inserindo Trecho 1 na Seção 5.3...")
insert_code_box_before(p_5_4_header, 
                       'Quadro 1 – Rotina de Ingestão e Mapeamento de Arquivos CSV de Atletas (handleImportAtletasCSV)', 
                       code_1, 
                       'App.jsx, linhas 353–378')

print("Inserindo Trecho 2 na Seção 5.4...")
insert_code_box_before(p_5_5_header, 
                       'Quadro 2 – Rotina de Filtragem de Candidatos Elegíveis e Aplicação de Bloqueios (RF06)', 
                       code_2, 
                       'App.jsx, linhas 154–177')

print("Inserindo Trecho 3 na Seção 5.5...")
insert_code_box_before(p_5_5_result_ph, 
                       'Quadro 3 – Algoritmo de Validação Pré-Voo em Quatro Dimensões (getValidationStatus)', 
                       code_3, 
                       'App.jsx, linhas 274–293')

print("Inserindo Trecho 4 na Seção 5.5...")
insert_code_box_before(p_5_6_header, 
                       'Quadro 4 – Cálculo da Penalização Etária e Pontuação da Guarnição (getAdjustedScore e b_jk)', 
                       code_4, 
                       'App.jsx, linhas 180–225')

doc.save(DOC_PATH)
print("Sucesso! Todos os 4 trechos de código foram inseridos com caixas estilizadas e legendas.")
